from pathlib import Path
from typing import List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from models import RevenueResult
from utils.logger import logger
from utils.formatters import format_number_bracket, format_variance_pct
from config import settings


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

HEADER_FILL = RGBColor(0xFF, 0xFF, 0xA3)
DARK_TEXT = RGBColor(0x0B, 0x30, 0x41)
SUBTOTAL_FILL = RGBColor(0x4A, 0x6F, 0x8A)
SUBTOTAL_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
LT_BLUE = RGBColor(0x20, 0x38, 0x4D)

FONT_NAME = "Arial"
HEADER_FONT_SIZE = Pt(10)
DATA_FONT_SIZE = Pt(9)


def generate_pptx(
    revenue_data: List[RevenueResult],
    report_month: str,
    report_year: int,
    output_dir: str,
) -> str:
    logger.info(f"Generating PPTX for {report_month} {report_year}")

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    _add_title(slide)
    _add_unit_label(slide)

    num_rows = len(revenue_data) + 2
    num_cols = 11

    table_shape = slide.shapes.add_table(
        num_rows,
        num_cols,
        Inches(0.3),
        Inches(1.1),
        Inches(12.7),
        Inches(5.5),
    )
    table = table_shape.table

    table.columns[0].width = Inches(2.5)
    for i in range(1, 11):
        table.columns[i].width = Inches(1.0)

    _set_header_rows(table, report_month, report_year)
    _populate_data_rows(table, revenue_data)
    _style_total_rows(table, revenue_data)

    output_filename = f"Revenue_{report_month}_{report_year}.pptx"
    output_path = Path(output_dir) / output_filename
    prs.save(str(output_path))

    logger.info(f"PPTX saved to: {output_path}")
    return str(output_path)


def _add_title(slide):
    txBox = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9), Inches(0.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = "Summary \u2013 Revenue"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = DARK_TEXT
    p.font.name = FONT_NAME


def _add_unit_label(slide):
    txBox = slide.shapes.add_textbox(
        Inches(10.84), Inches(0.75), Inches(1.41), Inches(0.30)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = "< Rs. Mn >"
    run.font.name = FONT_NAME
    run.font.size = Pt(12)
    run.font.color.rgb = DARK_TEXT


def _set_header_rows(table, report_month: str, report_year: int):
    short_month = report_month[:3] if report_month else "Mon"
    py_year = report_year - 1

    table.cell(0, 0).merge(table.cell(1, 0))
    _set_cell(
        table,
        0,
        0,
        "",
        bold=True,
        fill_color=HEADER_FILL,
        font_color=DARK_TEXT,
        font_size=HEADER_FONT_SIZE,
    )

    table.cell(0, 1).merge(table.cell(0, 4))
    _set_cell(
        table,
        0,
        1,
        f"Month \u2013 {short_month} '{str(report_year)[-2:]}",
        bold=True,
        fill_color=HEADER_FILL,
        font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER,
        font_size=HEADER_FONT_SIZE,
    )

    table.cell(0, 5).merge(table.cell(1, 5))
    _set_cell(
        table,
        0,
        5,
        f"{py_year}\nMonth",
        bold=True,
        fill_color=HEADER_FILL,
        font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER,
        font_size=HEADER_FONT_SIZE,
    )

    table.cell(0, 6).merge(table.cell(0, 9))
    _set_cell(
        table,
        0,
        6,
        f"YTD {short_month} '{str(report_year)[-2:]}",
        bold=True,
        fill_color=HEADER_FILL,
        font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER,
        font_size=HEADER_FONT_SIZE,
    )

    table.cell(0, 10).merge(table.cell(1, 10))
    _set_cell(
        table,
        0,
        10,
        f"{py_year}\nYTD",
        bold=True,
        fill_color=HEADER_FILL,
        font_color=DARK_TEXT,
        alignment=PP_ALIGN.CENTER,
        font_size=HEADER_FONT_SIZE,
    )

    for ci, txt in [
        (1, "Act"),
        (2, "Bud"),
        (3, "Vari"),
        (4, "Vari%"),
        (6, "Act"),
        (7, "Bud"),
        (8, "Vari"),
        (9, "Vari%"),
    ]:
        _set_cell(
            table,
            1,
            ci,
            txt,
            bold=True,
            fill_color=HEADER_FILL,
            font_color=DARK_TEXT,
            alignment=PP_ALIGN.CENTER,
            font_size=HEADER_FONT_SIZE,
        )

    for r in range(2):
        for c in range(11):
            _set_cell_fill(table.cell(r, c), HEADER_FILL)


def _populate_data_rows(table, revenue_data: List[RevenueResult]):
    for row_idx, data in enumerate(revenue_data):
        excel_row = row_idx + 2

        label = data.category
        _set_cell(
            table,
            excel_row,
            0,
            label,
            bold=data.is_subtotal,
            font_color=DARK_TEXT,
            alignment=PP_ALIGN.LEFT,
        )

        _set_cell(
            table,
            excel_row,
            1,
            format_number_bracket(data.month_actual),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            2,
            format_number_bracket(data.month_budget),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            3,
            format_number_bracket(data.month_variance),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            4,
            _format_var_pct_display(data.month_variance_pct),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            5,
            format_number_bracket(data.py_month_actual),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            6,
            format_number_bracket(data.ytd_actual),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            7,
            format_number_bracket(data.ytd_budget),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            8,
            format_number_bracket(data.ytd_variance),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            9,
            _format_var_pct_display(data.ytd_variance_pct),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )
        _set_cell(
            table,
            excel_row,
            10,
            format_number_bracket(data.py_ytd_actual),
            bold=data.is_subtotal,
            alignment=PP_ALIGN.RIGHT,
        )

        is_tot = data.is_subtotal
        for c in range(11):
            cell = table.cell(excel_row, c)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.bold = is_tot
                if data.is_grand_total:
                    paragraph.font.color.rgb = SUBTOTAL_TEXT
                else:
                    paragraph.font.color.rgb = DARK_TEXT
            if data.is_grand_total:
                _set_cell_fill(cell, SUBTOTAL_FILL)
            elif is_tot:
                _set_cell_fill(cell, SUBTOTAL_FILL)


def _style_total_rows(table, revenue_data: List[RevenueResult]):
    for row_idx, data in enumerate(revenue_data):
        if data.is_subtotal:
            excel_row = row_idx + 2
            for c in range(11):
                cell = table.cell(excel_row, c)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.bold = True
                    if data.is_grand_total:
                        paragraph.font.color.rgb = SUBTOTAL_TEXT
                    else:
                        paragraph.font.color.rgb = DARK_TEXT
                if data.is_grand_total:
                    _set_cell_fill(cell, SUBTOTAL_FILL)


def _set_cell(
    table,
    row,
    col,
    text,
    bold=False,
    font_color=None,
    fill_color=None,
    alignment=PP_ALIGN.LEFT,
    font_size=None,
):
    cell = table.cell(row, col)
    cell.text = ""

    if text is not None and text != "":
        p = cell.text_frame.paragraphs[0]
        p.alignment = alignment
        run = p.add_run()
        run.text = str(text)
        run.font.name = FONT_NAME
        run.font.size = font_size or DATA_FONT_SIZE
        run.font.bold = bold
        if font_color:
            run.font.color.rgb = font_color

    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Emu(0)
    cell.margin_right = Emu(0)
    cell.margin_top = Emu(0)
    cell.margin_bottom = Emu(0)

    if fill_color:
        _set_cell_fill(cell, fill_color)


def _set_cell_fill(cell, color: RGBColor):
    from pptx.oxml.ns import qn

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("a:solidFill")):
        tcPr.remove(existing)
    solidFill = tcPr.makeelement(qn("a:solidFill"), {})
    srgbClr = solidFill.makeelement(qn("a:srgbClr"), {"val": f"{color}"})
    solidFill.append(srgbClr)
    tcPr.append(solidFill)


def _format_var_pct_display(value: Optional[float]) -> str:
    if value is None:
        return "#DIV/0!"
    if abs(value) >= 1000:
        return f"{value:,.0f}%"
    if value == int(value):
        return f"{int(value)}%"
    return f"{value:.2f}%"
