import os
from pathlib import Path
from typing import List, Optional
from utils.logger import logger
from config import settings


class ValidationResult:
    def __init__(self):
        self.errors: List[str] = []
        self.warnings: List[str] = []
        self.is_valid: bool = True

    def add_error(self, message: str):
        self.errors.append(message)
        self.is_valid = False
        logger.error(f"Validation error: {message}")

    def add_warning(self, message: str):
        self.warnings.append(message)
        logger.warning(f"Validation warning: {message}")


def validate_uploaded_files(
    tb_current_path: Optional[str],
    tb_previous_path: Optional[str],
    budget_path: Optional[str],
    mapping_path: Optional[str],
) -> ValidationResult:
    result = ValidationResult()

    if not tb_current_path or not os.path.exists(tb_current_path):
        result.add_error("Current Year Trial Balance file is missing or not found.")
    elif not _is_valid_extension(tb_current_path):
        result.add_error("Current Year Trial Balance must be an Excel file (.xlsx/.xls).")

    if not tb_previous_path or not os.path.exists(tb_previous_path):
        result.add_error("Previous Year Trial Balance file is missing or not found.")
    elif not _is_valid_extension(tb_previous_path):
        result.add_error("Previous Year Trial Balance must be an Excel file (.xlsx/.xls).")

    if not budget_path or not os.path.exists(budget_path):
        result.add_error("Revenue Budget Workbook is missing or not found.")
    elif not _is_valid_extension(budget_path):
        result.add_error("Revenue Budget Workbook must be an Excel file (.xlsx/.xls).")

    if not mapping_path or not os.path.exists(mapping_path):
        result.add_error("Revenue Mapping Workbook is missing or not found.")
    elif not _is_valid_extension(mapping_path):
        result.add_error("Revenue Mapping Workbook must be an Excel file (.xlsx/.xls).")

    return result


def validate_mapping_workbook(mapping_path: str) -> ValidationResult:
    import openpyxl
    result = ValidationResult()
    try:
        wb = openpyxl.load_workbook(mapping_path, read_only=True, data_only=True)
    except Exception as e:
        result.add_error(f"Cannot open mapping workbook: {e}")
        return result

    if "Code Mapping" not in wb.sheetnames:
        result.add_error("Mapping workbook missing 'Code Mapping' sheet.")
        wb.close()
        return result

    ws = wb["Code Mapping"]
    header_row = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if row[0] and "heading" in str(row[0]).lower():
            header_row = row_idx
            break

    if header_row is None:
        result.add_warning("Could not find header row in Code Mapping sheet. Using row 3 as default.")
        header_row = 3

    wb.close()
    logger.info(f"Mapping workbook validated. Header at row {header_row}")
    return result


def validate_budget_workbook(budget_path: str) -> ValidationResult:
    import openpyxl
    result = ValidationResult()
    try:
        wb = openpyxl.load_workbook(budget_path, read_only=True, data_only=True)
    except Exception as e:
        result.add_error(f"Cannot open budget workbook: {e}")
        return result

    frm_found = False
    for name in wb.sheetnames:
        if name.strip().upper() == "FRM":
            frm_found = True
            break

    if not frm_found:
        result.add_error("Budget workbook missing 'FRM' sheet.")
    else:
        logger.info("Budget workbook validated. FRM sheet found.")

    wb.close()
    return result


def validate_trial_balance(tb_path: str, label: str) -> ValidationResult:
    import openpyxl
    result = ValidationResult()
    try:
        wb = openpyxl.load_workbook(tb_path, read_only=True, data_only=True)
    except Exception as e:
        result.add_error(f"Cannot open {label} Trial Balance: {e}")
        return result

    if len(wb.sheetnames) == 0:
        result.add_error(f"{label} Trial Balance has no sheets.")
        wb.close()
        return result

    ws = wb[wb.sheetnames[0]]
    has_data = False
    for row in ws.iter_rows(min_row=1, max_row=50, values_only=True):
        for cell in row:
            if cell is not None and str(cell).strip():
                has_data = True
                break
        if has_data:
            break

    if not has_data:
        result.add_error(f"{label} Trial Balance appears to be empty.")

    wb.close()
    logger.info(f"{label} Trial Balance validated. Sheet: {wb.sheetnames[0]}")
    return result


def _is_valid_extension(file_path: str) -> bool:
    ext = Path(file_path).suffix.lower()
    return ext in settings.ALLOWED_EXTENSIONS
